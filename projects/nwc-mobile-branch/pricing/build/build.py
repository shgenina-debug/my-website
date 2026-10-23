import sys
sys.path.insert(0,__import__('os').path.dirname(__import__('os').path.abspath(__file__)))
from data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

wb=Workbook()
AR="Arial"
F=lambda **k: Font(name=AR, **k)
BLUE=F(color="0000FF"); GREEN=F(color="008000"); BOLD=F(bold=True); HDR=F(bold=True,color="FFFFFF")
HFILL=PatternFill("solid",fgColor="1F4E78"); YEL=PatternFill("solid",fgColor="FFFF00"); SUB=PatternFill("solid",fgColor="DDEBF7"); TOT=PatternFill("solid",fgColor="F2F2F2")
thin=Side(style="thin",color="BFBFBF"); BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
WRAP=Alignment(wrap_text=True,vertical="top")
NUM='#,##0.00;(#,##0.00);"-"'; NUM0='#,##0;(#,##0);"-"'; PCT='0.0%;(0.0%);"-"'

def sheet(title):
    ws=wb.create_sheet(title); ws.sheet_view.rightToLeft=True; return ws
def header(ws,row,labels,widths=None):
    for i,l in enumerate(labels,1):
        c=ws.cell(row,i,l); c.font=HDR; c.fill=HFILL; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); c.border=BOX
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name!=AR:
                c.font=Font(name=AR,bold=c.font.bold if c.font else False,color=c.font.color if c.font else None)

wb.remove(wb.active)
# ---------- Inputs ----------
inp=sheet("المدخلات")
inp["A1"]="المدخلات والافتراضات — فرع خدمة العملاء المتنقل"; inp["A1"].font=F(bold=True,size=14)
inp["A2"]="الخلايا الزرقاء مدخلات تُعدَّل، والصفراء افتراضات رئيسية. كل الأسعار قبل ضريبة القيمة المضافة. تاريخ مراجعة الأسعار: 21/10/2026."
header(inp,3,["البند","القيمة","ملاحظة / المصدر"],[44,14,80])
rows=[
("SAR",1,"الريال السعودي"),
("USD",3.75,"سعر الربط الرسمي"),
("EUR",4.35,"افتراض — يُحدَّث بسعر يوم التسعير"),
("GBP",5.05,"افتراض — يُحدَّث بسعر يوم التسعير"),
("AED","=3.75/3.6725","مشتق من ربط الريال والدرهم بالدولار"),
("AUD",2.45,"افتراض — يُحدَّث بسعر يوم التسعير"),
]
r=4
for k,v,n in rows:
    inp.cell(r,1,k); c=inp.cell(r,2,v); c.font=BLUE; c.number_format='0.0000'; inp.cell(r,3,n); r+=1
FX="المدخلات!$A$4:$B$9"
inp.cell(11,1,"معامل تكلفة الاستيراد (شحن + جمارك 5% + تخليص وتأمين)"); c=inp.cell(11,2,1.15); c.font=BLUE; c.fill=YEL; inp.cell(11,3,"افتراض: شحن ≈ 6% + رسوم جمركية 5% + تخليص وتأمين ونقل داخلي ≈ 4%")
inp.cell(13,1,"تكلفة يوم العامل الماهر (شاملة السكن والمقابل المالي والتأمينات)"); c=inp.cell(13,2,380); c.font=BLUE; c.fill=YEL
inp.cell(13,3,"افتراض: راتب لحام/فني 6,000–12,000 ر.س شهرياً (Paylab) + تكاليف نظامية ÷ 26 يوم عمل")
inp.cell(14,1,"تكلفة يوم العامل المساعد"); c=inp.cell(14,2,250); c.font=BLUE; c.fill=YEL; inp.cell(14,3,"افتراض")
inp.cell(16,1,"مدة المشروع (أشهر)"); c=inp.cell(16,2,3); c.font=BLUE; inp.cell(16,3,"التسليم خلال 90 يوماً")
mk=[("المصاريف العامة للشركة (من تكلفة المشروع)",0.05),("تكلفة التمويل (من تكلفة المشروع)",0.025),("الضمانات البنكية والتأمين (من تكلفة المشروع)",0.01),
    ("المخاطر والطوارئ (من تكلفة المشروع)",0.05),("الربح (من إجمالي التكلفة)",0.10),("ضريبة القيمة المضافة",0.15)]
for i,(k,v) in enumerate(mk):
    inp.cell(18+i,1,k); c=inp.cell(18+i,2,v); c.font=BLUE; c.number_format='0.0%'
    if i<5: c.fill=YEL
inp.cell(19,3,"بافتراض عدم وجود دفعة مقدمة؛ يُخفَّض إن نصت الكراسة عليها")
inp.cell(25,1,"مستويات الثقة في الأسعار"); inp.cell(25,1).font=BOLD
leg=[(S,"سعر من متجر أو جهة سعودية"),(I,"سعر من متجر دولي + معامل الاستيراد"),(E,"تقدير مبني على أكثر من مصدر أو منتج مقارب"),(A,"تقدير هندسي — يجب استبداله بعرض سعر مورد")]
for i,(k,n) in enumerate(leg): inp.cell(26+i,1,k); inp.cell(26+i,3,n)
LAB={"LAB_S":"=المدخلات!$B$13","LAB_H":"=المدخلات!$B$14"}

# ---------- Breakdown ----------
bd=sheet("تكسير البنود")
bd["A1"]="تكسير البنود — التكلفة المباشرة (قبل الضريبة)"; bd["A1"].font=F(bold=True,size=14)
cols=["البند","م","الوصف","الوحدة","الكمية","سعر الوحدة (بالعملة)","العملة","مستورد؟","سعر الوحدة (ر.س)","الإجمالي (ر.س)","النوع","مستوى الثقة","الأساس","الرابط"]
header(bd,3,cols,[7,5,52,8,8,13,7,8,13,14,13,22,48,40])
r=4; first=r; seq={}
order=list(ITEMS)
for code in order:
    lines=[x for x in R if x[0]==code]
    t=bd.cell(r,3,f"البند {code} — {ITEMS[code]}"); t.font=BOLD
    for cc in range(1,15): bd.cell(r,cc).fill=SUB
    r+=1
    for n,(it,desc,unit,qty,rate,cur,imp,kind,conf,basis,url) in enumerate(lines,1):
        bd.cell(r,1,code); bd.cell(r,2,n); bd.cell(r,3,desc).alignment=WRAP; bd.cell(r,4,unit)
        c=bd.cell(r,5,qty); c.font=BLUE; c.number_format='#,##0.##'
        if isinstance(rate,str): c=bd.cell(r,6,LAB[rate]); c.font=GREEN
        else: c=bd.cell(r,6,rate); c.font=BLUE
        c.number_format=NUM
        bd.cell(r,7,"ر.س" if cur=="SAR" else cur); bd.cell(r,8,"نعم" if imp else "لا")
        cur_cell=f'IF(G{r}="ر.س","SAR",G{r})'
        bd.cell(r,9,f'=F{r}*VLOOKUP({cur_cell},{FX},2,FALSE)*IF(H{r}="نعم",المدخلات!$B$11,1)').number_format=NUM
        bd.cell(r,10,f"=E{r}*I{r}").number_format=NUM
        bd.cell(r,11,kind); bd.cell(r,12,conf); bd.cell(r,13,basis).alignment=WRAP
        if url: bd.cell(r,14,url).hyperlink=url
        r+=1
    s=bd.cell(r,3,f"إجمالي البند {code}"); s.font=BOLD
    c=bd.cell(r,10,f'=SUMIF($A${first}:$A${r-1},"{code}",$J${first}:$J${r-1})'); c.font=BOLD; c.number_format=NUM
    for cc in range(1,15): bd.cell(r,cc).fill=TOT
    r+=2
last=r
bd.cell(r,3,"إجمالي التكلفة المباشرة").font=BOLD
c=bd.cell(r,10,f'=SUMPRODUCT(($A${first}:$A${last-1}<>"")*$J${first}:$J${last-1})'); c.font=BOLD; c.number_format=NUM
BDR=f"'تكسير البنود'!$A${first}:$A${last}"; BDJ=f"'تكسير البنود'!$J${first}:$J${last}"; BDL=f"'تكسير البنود'!$L${first}:$L${last}"; BDK=f"'تكسير البنود'!$K${first}:$K${last}"
bd.freeze_panes="D4"; bd.auto_filter.ref=f"A3:N{last-1}"

# ---------- Overheads ----------
oh=sheet("مصاريف المشروع")
oh["A1"]="مصاريف المشروع غير المباشرة"; oh["A1"].font=F(bold=True,size=14)
header(oh,3,["البند","العدد","التكلفة الشهرية (ر.س)","المدة (شهر)","مبلغ مقطوع (ر.س)","الإجمالي (ر.س)"],[44,8,18,12,18,16])
r=4
for name,cnt,mon,lump in OH:
    oh.cell(r,1,name+" — افتراض")
    if lump is None:
        for col,v in ((2,cnt),(3,mon)):
            c=oh.cell(r,col,v); c.font=BLUE; c.number_format=NUM0
        c=oh.cell(r,4,"=المدخلات!$B$16"); c.font=GREEN
    else:
        c=oh.cell(r,5,lump); c.font=BLUE; c.number_format=NUM0
    oh.cell(r,6,f'=IF(E{r}<>"",E{r},B{r}*C{r}*D{r})').number_format=NUM0
    r+=1
oh.cell(r,1,"الإجمالي").font=BOLD; c=oh.cell(r,6,f"=SUM(F4:F{r-1})"); c.font=BOLD; c.number_format=NUM0
OHT=f"'مصاريف المشروع'!$F${r}"

# ---------- Offer summary ----------
sm=sheet("ملخص العرض")
sm["A1"]="ملخص العرض المالي"; sm["A1"].font=F(bold=True,size=14)
header(sm,3,["المكوّن","المبلغ (ر.س)","من سعر العرض"],[46,18,16])
SL=[("التكلفة المباشرة","=SUM('ملخص البنود'!C5:C22)"),
("مصاريف المشروع غير المباشرة",f"={OHT}"),
("تكلفة المشروع","=B4+B5"),
("المصاريف العامة للشركة","=B6*المدخلات!$B$18"),
("تكلفة التمويل","=B6*المدخلات!$B$19"),
("الضمانات البنكية والتأمين","=B6*المدخلات!$B$20"),
("المخاطر والطوارئ","=B6*المدخلات!$B$21"),
("إجمالي التكلفة","=SUM(B6:B10)"),
("الربح","=B11*المدخلات!$B$22"),
("سعر العرض قبل الضريبة (محسوب)","=B11+B12"),
("معامل التحميل على التكلفة المباشرة","=B13/B4"),
("مجموع جدول الكميات المسعّر (المعتمد في العرض)","='ملخص البنود'!I23"),
("ضريبة القيمة المضافة","=B15*المدخلات!$B$23"),
("الإجمالي شامل الضريبة","=B15+B16"),
("",""),
("سعر العرض السابق قبل الضريبة (21/10/2026)",1299139.15),
("الفرق عن العرض السابق","=B15-B19"),
("نسبة الفرق","=IF(B19=0,0,B20/B19)"),
]
for i,(k,f) in enumerate(SL):
    rr=4+i; sm.cell(rr,1,k)
    if f!="":
        c=sm.cell(rr,2,f); c.number_format=NUM
        if k.startswith("معامل"): c.number_format='0.0000'
        if k.startswith("نسبة"): c.number_format=PCT
        if k.startswith("سعر العرض السابق"): c.font=BLUE
    if rr in (4,5,6,7,8,9,10,11,12,13): sm.cell(rr,3,f"=IF($B$13=0,0,B{rr}/$B$13)").number_format=PCT
for rr in (6,11,13,15,17): sm.cell(rr,1).font=BOLD; sm.cell(rr,2).font=BOLD
sm.cell(23,1,"تكوين التكلفة المباشرة حسب مستوى الثقة").font=BOLD
header(sm,24,["مستوى الثقة","المبلغ (ر.س)","من التكلفة المباشرة"])
for i,k in enumerate([S,I,E,A]):
    rr=25+i; sm.cell(rr,1,k); sm.cell(rr,2,f'=SUMIF({BDL},A{rr},{BDJ})').number_format=NUM; sm.cell(rr,3,f"=IF($B$4=0,0,B{rr}/$B$4)").number_format=PCT
sm.cell(30,1,"تكوين التكلفة المباشرة حسب النوع").font=BOLD
header(sm,31,["النوع","المبلغ (ر.س)","من التكلفة المباشرة"])
for i,k in enumerate([M,L,SC]):
    rr=32+i; sm.cell(rr,1,k); sm.cell(rr,2,f'=SUMIF({BDK},A{rr},{BDJ})').number_format=NUM; sm.cell(rr,3,f"=IF($B$4=0,0,B{rr}/$B$4)").number_format=PCT
sm.cell(36,1,"تحقق: مجموع التكوين يساوي التكلفة المباشرة"); sm.cell(36,2,"=ROUND(SUM(B25:B28)-B4,2)").number_format=NUM

# ---------- Item summary / priced BOQ ----------
it=sheet("ملخص البنود")
it["A1"]="ملخص البنود وجدول الكميات المسعّر"; it["A1"].font=F(bold=True,size=14)
it["A2"]="التكلفة المباشرة من ورقة التكسير، والسعر = التكلفة × معامل التحميل (توزيع متوازن). الأسعار قبل الضريبة."
header(it,4,["البند","الوصف","التكلفة المباشرة (ر.س)","تقدير التصميم R03 (ر.س)","الفرق (ر.س)","الفرق %","منها افتراضات (ر.س)","نسبة الافتراضات","سعر البند في العرض (ر.س)"],[7,40,18,18,14,10,18,12,20])
for i,code in enumerate(order):
    rr=5+i
    it.cell(rr,1,code); it.cell(rr,2,ITEMS[code])
    it.cell(rr,3,f'=SUMIF({BDR},A{rr},{BDJ})').number_format=NUM
    c=it.cell(rr,4,OLD[code]); c.font=BLUE; c.number_format=NUM0
    it.cell(rr,5,f"=C{rr}-D{rr}").number_format=NUM0
    it.cell(rr,6,f"=IF(D{rr}=0,0,E{rr}/D{rr})").number_format=PCT
    it.cell(rr,7,f'=SUMIFS({BDJ},{BDR},A{rr},{BDL},"{A}")').number_format=NUM0
    it.cell(rr,8,f"=IF(C{rr}=0,0,G{rr}/C{rr})").number_format=PCT
    it.cell(rr,9,f"=ROUND(C{rr}*'ملخص العرض'!$B$14,2)").number_format=NUM
it.cell(23,2,"الإجمالي").font=BOLD
for col in "CDEGI":
    c=it[f"{col}23"]; c.value=f"=SUM({col}5:{col}22)"; c.font=BOLD; c.number_format=NUM
it["F23"]="=IF(D23=0,0,E23/D23)"; it["F23"].number_format=PCT; it["H23"]="=IF(C23=0,0,G23/C23)"; it["H23"].number_format=PCT
it["B24"]="ضريبة القيمة المضافة"; it["I24"]="=I23*المدخلات!$B$23"; it["I24"].number_format=NUM
it["B25"]="الإجمالي شامل الضريبة"; it["B25"].font=BOLD; it["I25"]="=I23+I24"; it["I25"].number_format=NUM; it["I25"].font=BOLD
it["B26"]="فرق التقريب عن الملخص"; it["I26"]="=I23-'ملخص العرض'!B13"; it["I26"].number_format=NUM
it.freeze_panes="C5"

# ---------- Optional items ----------
op=sheet("البنود الاختيارية")
op["A1"]="البنود الاختيارية (خارج إجمالي العرض)"; op["A1"].font=F(bold=True,size=14)
op["A2"]="معامل الإضافات = (1 + عامة + تمويل + ضمانات + طوارئ) × (1 + ربح)، دون مصاريف المشروع لأنها مغطاة في العرض الأساسي."
op["A3"]="معامل البنود الاختيارية"; op["B3"]="=(1+المدخلات!$B$18+المدخلات!$B$19+المدخلات!$B$20+المدخلات!$B$21)*(1+المدخلات!$B$22)"; op["B3"].number_format='0.0000'
header(op,5,["البند","الوصف","الوحدة","التكلفة المباشرة (ر.س)","السعر في العرض (ر.س)","الأساس"],[8,52,8,18,18,70])
OPT=[
("خ-1","استبدال المولد بمولد صامت 10 ك.ف.أ (الفرق فقط)","مقطوع","=(4800-3112.5)*VLOOKUP(\"EUR\","+FX+",2,FALSE)*المدخلات!$B$11+2500","مولد 10 ك.ف.أ من الفئة نفسها ≈ 4,800 يورو (افتراض) ناقص المولد الأساسي 3,112.50 يورو، + 2,500 ر.س تعديل الحجرة والكابلات"),
("خ-2","مكيف سقفي RV انفرتر إضافي لمنطقة الموظفين","عدد","=(1900+250)*VLOOKUP(\"USD\","+FX+",2,FALSE)*المدخلات!$B$11+2*المدخلات!$B$13+600","سعر المكيف وصندوق التوزيع نفسه في البند 06 + يوما تركيب + 600 ر.س تمديدات وقاطع"),
("خ-3","عزل حراري للأرضية","—",0,"مشمول في البند 02 (صوف صخري 50 مم)"),
("خ-4أ","لابتوب Dell Latitude 5450 Ultra 7 16/512 مع التهيئة","عدد","=4296+150","ابتكار 4,940 ر.س شامل الضريبة + 150 ر.س تهيئة"),
("خ-4ب","تابلت Samsung Galaxy Tab S11 5G","عدد","=1049/1.2*VLOOKUP(\"EUR\","+FX+",2,FALSE)","سعر الإطلاق الأوروبي 1,049 يورو شامل ضريبة ≈ 20% — يُؤكد من جرير أو إكسترا"),
("خ-4ج","نقطة وصول Cisco CW9176I مع التركيب والتهيئة (دون تراخيص الاشتراك)","عدد","=1505.37*VLOOKUP(\"USD\","+FX+",2,FALSE)*المدخلات!$B$11+300","Hummingbird Networks 1,505.37 دولار + معامل الاستيراد + 300 ر.س تركيب"),
]
for i,(code,desc,unit,f,basis) in enumerate(OPT):
    rr=6+i; op.cell(rr,1,code); op.cell(rr,2,desc).alignment=WRAP; op.cell(rr,3,unit)
    op.cell(rr,4,f).number_format=NUM0
    op.cell(rr,5,f"=ROUND(D{rr}*$B$3,0)").number_format=NUM0
    op.cell(rr,6,basis).alignment=WRAP

# ---------- Sources ----------
so=sheet("المصادر")
so["A1"]="مصادر الأسعار (تمت مراجعتها في 21/10/2026)"; so["A1"].font=F(bold=True,size=14)
header(so,3,["البند","المكوّن","المصدر","الرابط"],[7,50,50,70])
rr=4
for x in R:
    if x[10]:
        so.cell(rr,1,x[0]); so.cell(rr,2,x[1]).alignment=WRAP; so.cell(rr,3,x[9]).alignment=WRAP; c=so.cell(rr,4,x[10]); c.hyperlink=x[10]; rr+=1
so.cell(rr+1,1,"ملاحظة: لم تُفتح صفحات المواقع الحكومية والمتاجر مباشرة لأن الشبكة تحجبها؛ الأسعار مأخوذة من نتائج محركات البحث وتُؤكد بعروض موردين قبل التسليم.")

wb._sheets=[wb[n] for n in ["ملخص العرض","ملخص البنود","المدخلات","تكسير البنود","مصاريف المشروع","البنود الاختيارية","المصادر"]]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            f=c.font
            c.font=Font(name=AR,bold=f.bold,size=f.size,color=f.color)
import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","mobile_branch_pricing.xlsx")
wb.save(out); print(out, [w.title for w in wb.worksheets])
